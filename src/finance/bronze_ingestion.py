"""
bronze_ingestion.py — Couche Bronze
====================================
Parseur adapté à la structure réelle des fichiers Budget_YYYY.xlsx.

Structure Excel détectée :
  - Colonne A : libellé (poste, section, total...)
  - Colonne B : montant (float, formule Excel, ou None)
  - Chaque feuille = un mois
  - Sections identifiées par des mots-clés (sans montant associé)
  - Lignes ignorées : sous-totaux, totaux, en-têtes

Sortie (Parquet Bronze) :
  annee | mois | section | libelle | montant
"""

import re
import openpyxl
import pandas as pd
from pathlib import Path

from src.shared.config import BRONZE_PATH
from src.shared.utils import ensure_directory_exists


# ─── Mots-clés ────────────────────────────────────────────────────────────────

# Lignes qui marquent le début d'une nouvelle section (pas un poste réel)
SECTION_KEYWORDS = [
    "revenus fixes", "revenus variables",
    "dépenses essentielles", "dépenses optionnelles",
    "revenus", "dépenses",
]

# Lignes à ignorer complètement (totaux, récap, semaines...)
IGNORE_KEYWORDS = [
    "sous-total", "total", "récapitulatif",
    "revenu - dépenses", "semaine",
]

# Lignes d'épargne → redirigées vers section "Épargne" dédiée
EPARGNE_KEYWORDS = [
    "épargne", "epargne",
]


# ─── Helpers ──────────────────────────────────────────────────────────────────

def is_section_header(label: str) -> bool:
    """Retourne True si la ligne est un titre de section (pas un poste réel)."""
    low = label.strip().lower()
    return any(low.startswith(kw) for kw in SECTION_KEYWORDS)


def is_ignored(label: str) -> bool:
    """Retourne True si la ligne est un total / récap à ne pas ingérer."""
    low = label.strip().lower()
    return any(kw in low for kw in IGNORE_KEYWORDS)


def is_epargne(label: str) -> bool:
    """Retourne True si la ligne concerne l'épargne → section dédiée."""
    low = label.strip().lower()
    return any(kw in low for kw in EPARGNE_KEYWORDS)


def is_formula(value) -> bool:
    """Retourne True si la valeur est une formule Excel non résolue."""
    return isinstance(value, str) and value.startswith("=")


# ─── Parseur d'une feuille ────────────────────────────────────────────────────

def parse_sheet(ws, annee: int, mois: str) -> list[dict]:
    """
    Parcourt une feuille Excel et retourne une liste de postes budgétaires.
    Chaque dict : {annee, mois, section, libelle, montant}
    """
    records = []
    current_section = "Non catégorisé"

    for row in ws.iter_rows(values_only=True):
        label = row[0]
        montant = row[1] if len(row) > 1 else None

        if not label:
            continue

        label = str(label).strip()

        # Ignorer l'en-tête du tableau (ligne "Revenus / Montant")
        if label.lower() == "revenus" and str(montant).lower() == "montant":
            continue

        # Ignorer totaux / récap
        if is_ignored(label):
            continue

        # Détecter une nouvelle section
        if is_section_header(label):
            current_section = label
            continue

        # Ignorer les formules Excel non résolues
        if is_formula(montant):
            continue

        # Ignorer les postes sans montant numérique
        if montant is None or not isinstance(montant, (int, float)):
            continue

        # Lignes d'épargne → section dédiée
        section = "Épargne" if is_epargne(label) else current_section

        records.append({
            "annee":   annee,
            "mois":    mois,
            "section": section,
            "libelle": label,
            "montant": float(montant),
        })

    return records


# ─── Parseur d'un fichier Excel ───────────────────────────────────────────────

def parse_workbook(path: Path) -> pd.DataFrame:
    """
    Charge toutes les feuilles d'un fichier Excel et retourne un DataFrame.
    Le nom du fichier doit contenir l'année (ex: Budget_2025.xlsx).
    """
    match = re.search(r"(\d{4})", path.name)
    if not match:
        raise ValueError(f"Impossible d'extraire l'année depuis : {path.name}")
    annee = int(match.group(1))

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    all_records = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        records = parse_sheet(ws, annee=annee, mois=sheet_name)
        all_records.extend(records)
        print(f"  ✅ [{sheet_name}] : {len(records)} postes extraits")

    wb.close()

    df = pd.DataFrame(all_records, columns=["annee", "mois", "section", "libelle", "montant"])
    return df


# ─── Validation ───────────────────────────────────────────────────────────────

def validate_bronze(df: pd.DataFrame) -> pd.DataFrame:
    """Vérifie que le DataFrame contient les colonnes attendues et des données."""
    required = ["annee", "mois", "section", "libelle", "montant"]
    missing = [c for c in required if c not in df.columns]
    if missing:
        raise ValueError(f"Colonnes manquantes : {missing}")
    if df.empty:
        raise ValueError("Aucune donnée extraite — vérifiez la structure des fichiers Excel.")
    if df["montant"].isna().all():
        raise ValueError("Tous les montants sont nuls — parsing probablement incorrect.")
    return df


# ─── Écriture Parquet ─────────────────────────────────────────────────────────

def write_to_bronze(df: pd.DataFrame, filename: str) -> None:
    """Écrit le DataFrame en Parquet dans le dossier Bronze."""
    ensure_directory_exists(BRONZE_PATH)
    output_path = Path(BRONZE_PATH) / filename
    df.to_parquet(output_path, index=False)
    print(f"  💾 Écrit : {output_path}  ({len(df)} lignes)")


# ─── Pipeline principal ───────────────────────────────────────────────────────

def run_bronze_ingestion(source: str, filename: str) -> pd.DataFrame:
    """Pipeline Bronze complet pour un fichier Excel."""
    path = Path(source)
    print(f"\n📥 Chargement : {path}")

    df = parse_workbook(path)

    print("🔍 Validation du schéma...")
    df = validate_bronze(df)

    print("📦 Écriture Bronze...")
    write_to_bronze(df, filename)

    print(f"✅ Terminé — {len(df)} postes ingérés")
    return df


# ─── Point d'entrée ───────────────────────────────────────────────────────────

if __name__ == "__main__":
    run_bronze_ingestion("inputs/Budget 2025.xlsx", "budget_2025.parquet")
    run_bronze_ingestion("inputs/Budget 2026.xlsx", "budget_2026.parquet")